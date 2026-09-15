"""Structural validation for course_overview/Machine_Learning_for_Neuroscience_Notebook_Overview.docx
(WP18 sec 6, sec 9).

python-docx is intentionally NOT part of this project's student-facing
``requirements.txt`` -- it is only used to generate/validate this one
administrative document, via the SYSTEM Python where it is already
installed. This test therefore skips itself (rather than failing) when
``docx`` is not importable in whatever interpreter runs it -- run it with an
interpreter that has python-docx to actually exercise it:

    python3 -m unittest tests.test_course_overview_docx -v
"""

from __future__ import annotations

import unittest
import zipfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT

    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False

REPO_ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = REPO_ROOT / "course_overview" / "Machine_Learning_for_Neuroscience_Notebook_Overview.docx"

EXPECTED_TITLE = "Machine Learning for Neuroscience — Practice Notebook Overview"
EXPECTED_SUBJECT_KEYWORDS = (
    "Exploratory Data Analysis",
    "Linear Regression",
    "KNN Regression",
    "Logistic Regression",
)
EXPECTED_COVERAGE_PHRASES = (
    "head(), tail(), and sample()",
    "honest held-out",
    "preselected k = 20",
    "fixed C = 1.0",
    "TN, FP, FN, TP",
)

# WP19: early lessons (Exercises 1-4) must not teach cross-validation or
# formal hyperparameter selection; the overview document must not claim it.
PROHIBITED_COVERAGE_PHRASES = (
    "cross-validation",
    "cross validation",
    "choosing k honestly",
    "choosing C honestly",
    "GridSearchCV",
)


@unittest.skipUnless(HAVE_DOCX, "python-docx not installed in this interpreter")
class CourseOverviewDocx(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not DOCX_PATH.exists():
            raise AssertionError(f"{DOCX_PATH} does not exist")
        cls.doc = Document(DOCX_PATH)

    def test_zip_structure_is_valid(self):
        with zipfile.ZipFile(DOCX_PATH) as zf:
            self.assertIsNone(zf.testzip())
            names = zf.namelist()
            self.assertIn("[Content_Types].xml", names)
            self.assertIn("word/document.xml", names)

    def test_title_is_present_and_exact(self):
        title_text = self.doc.paragraphs[0].text.strip()
        self.assertTrue(title_text)
        self.assertEqual(title_text, EXPECTED_TITLE)

    def test_exactly_one_table_five_rows_three_columns(self):
        self.assertEqual(len(self.doc.tables), 1)
        table = self.doc.tables[0]
        self.assertEqual(len(table.columns), 3)
        self.assertEqual(len(table.rows), 5)

    def test_header_row_is_exact(self):
        header = [c.text.strip() for c in self.doc.tables[0].rows[0].cells]
        self.assertEqual(header, ["Exercise", "Subject", "Covered materials"])

    def test_exercises_1_through_4_present_once_each(self):
        table = self.doc.tables[0]
        exercise_col = [table.rows[i].cells[0].text.strip() for i in range(1, len(table.rows))]
        self.assertEqual(exercise_col, ["Exercise 1", "Exercise 2", "Exercise 3", "Exercise 4"])

    def test_required_subjects_present(self):
        table = self.doc.tables[0]
        subjects = [table.rows[i].cells[1].text.strip() for i in range(1, len(table.rows))]
        for subject, keyword in zip(subjects, EXPECTED_SUBJECT_KEYWORDS):
            self.assertIn(keyword, subject)

    def test_key_coverage_phrases_present(self):
        table = self.doc.tables[0]
        covered = " ".join(table.rows[i].cells[2].text for i in range(1, len(table.rows)))
        for phrase in EXPECTED_COVERAGE_PHRASES:
            self.assertIn(phrase, covered)

    def test_no_early_cross_validation_or_tuning_claim(self):
        table = self.doc.tables[0]
        covered = " ".join(table.rows[i].cells[2].text for i in range(1, len(table.rows)))
        for phrase in PROHIBITED_COVERAGE_PHRASES:
            self.assertNotIn(phrase, covered)

    def test_rows_are_marked_cant_split(self):
        # WP19: fixes the WP18 layout defect where the Exercise 4 row split
        # awkwardly across a nearly empty second page.
        from docx.oxml.ns import qn

        table = self.doc.tables[0]
        for row in table.rows:
            trPr = row._tr.find(qn("w:trPr"))
            self.assertIsNotNone(trPr, "row is missing trPr")
            self.assertIsNotNone(trPr.find(qn("w:cantSplit")), "row is missing w:cantSplit")

    def test_document_title_metadata_is_set(self):
        self.assertTrue(self.doc.core_properties.title)

    def test_landscape_orientation(self):
        section = self.doc.sections[0]
        self.assertEqual(section.orientation, WD_ORIENT.LANDSCAPE)
        self.assertGreater(section.page_width, section.page_height)


if __name__ == "__main__":
    unittest.main()
