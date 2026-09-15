#!/usr/bin/env python3
"""Build the Word overview of Exercises 1-4 (WP18 sec 6).

python-docx is intentionally NOT added to this project's student-facing
``requirements.txt`` -- this is an administrative document, not part of the
student runtime. Run this script with an interpreter that already has
python-docx installed (e.g. the system ``python3``, not ``.venv``):

    python3 scripts/build_course_overview_docx.py

Structural validation lives in ``tests/test_course_overview_docx.py`` (also
run with an interpreter that has python-docx).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = REPO_ROOT / "course_overview" / "Machine_Learning_for_Neuroscience_Notebook_Overview.docx"

TITLE = "Machine Learning for Neuroscience — Practice Notebook Overview"
INTRO = (
    "This document summarizes the practice notebooks currently developed for "
    "the course, one row per exercise."
)

HEADER_BLUE = RGBColor(0x1F, 0x4E, 0x79)
HEADER_FILL_HEX = "1F4E79"

# Reconciled against the actual final notebooks (book/chapters/chapter_0{1..4})
# rather than relying solely on older WP descriptions.
ROWS = [
    (
        "Exercise 1",
        "Exploratory Data Analysis (EDA)",
        "ABIDE-II phenotypic table (13 curated columns); initial inspection with "
        "head(), tail(), and sample() (interactive comparison); dataset structure "
        "and data types; identifying hidden categorical variables; describe() "
        "summary statistics; quantifying missingness and a complete-case retention "
        "explorer (interactive); distributions and group comparisons; Pearson and "
        "Spearman correlations; interactive histogram and correlation explorer.",
    ),
    (
        "Exercise 2",
        "Linear Regression",
        "ABIDE neuroimaging ROI data (360 cortical-thickness predictors); age "
        "prediction with linear regression; StandardScaler inside a train/test "
        "split; honest held-out R² and MSE versus invalid resubstitution/leakage "
        "scoring; comparing feature sets (interactive activity); effect of sample "
        "size on model stability.",
    ),
    (
        "Exercise 3",
        "KNN Regression and the Bias–Variance Trade-off",
        "KNN regression for age prediction; scaling inside a Pipeline; choosing k "
        "honestly via training-only 5-fold cross-validation; honest versus invalid "
        "evaluation (interactive activity); the classic bias–variance tradeoff; "
        "empirical k = 1..N_fit curve; explore-k interactive activity.",
    ),
    (
        "Exercise 4",
        "Classification with Logistic Regression",
        "ABIDE autism-versus-control classification (360 ROI predictors); logistic "
        "regression and the sigmoid function; choosing C honestly via training-only "
        "cross-validation scored by ROC AUC; confusion matrix labelled with TN, FP, "
        "FN, TP; accuracy, sensitivity, specificity, ROC/AUC; decision-threshold "
        "interactive activity; class imbalance and misleading accuracy (interactive "
        "activity).",
    ),
]


def _set_cell_background(cell, color_hex: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_col_widths(table, widths_in) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            cell.width = Inches(width)


def build() -> None:
    doc = Document()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Machine Learning for Neuroscience course overview"

    title_par = doc.add_paragraph()
    title_par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_par.add_run(TITLE)
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = HEADER_BLUE

    intro_par = doc.add_paragraph(INTRO)
    intro_par.runs[0].font.size = Pt(12)
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for cell, text in zip(header_cells, ("Exercise", "Subject", "Covered materials")):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _set_cell_background(cell, HEADER_FILL_HEX)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for exercise, subject, covered in ROWS:
        row_cells = table.add_row().cells
        for cell, text in zip(row_cells, (exercise, subject, covered)):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(12)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    _set_col_widths(table, (1.0, 2.3, 8.2))

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(f"wrote {OUT_PATH.relative_to(REPO_ROOT)} ({OUT_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
